"""Вывод форматированного текста (Delta редактора Quill) в ячейку DOCX.

Поддерживается: жирный, курсив, подчёркивание, код в строке (моноширинный
шрифт с серым фоном), блоки кода, нумерованные и маркированные списки,
рисунки с подписью. Рисунки нумеруются по разделу уязвимости:
«Рисунок 4.1.1 – подпись», «Рисунок 4.1.2 – …» (сквозь PoC и результаты).
"""
import json
import re

from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from . import docx_helpers as H

CODE_FONT = "Courier New"
CODE_FILL = "F2F2F2"
_LEGACY_NUM = re.compile(r"^\s*Рисунок\s*[\d.]*\s*[–—-]?\s*", re.IGNORECASE)


def parse(value):
    """Delta из сохранённого значения или None для обычного текста (старые карточки)."""
    if isinstance(value, str) and value.startswith('{"ops"'):
        try:
            ops = json.loads(value).get("ops")
            if isinstance(ops, list):
                return ops
        except ValueError:
            pass
    return None


def figure_ids(value):
    return {op["insert"]["figure"].get("id") for op in (parse(value) or [])
            if isinstance(op.get("insert"), dict) and "figure" in op["insert"]}


def _shade(el_pr, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    el_pr.append(shd)


def _run(par, text, attrs, size, code_block=False):
    r = par.add_run(text)
    is_code = code_block or attrs.get("code")
    H._set_run(r, bold=bool(attrs.get("bold")), italic=bool(attrs.get("italic")),
               size=(size - 1.5 if is_code else size), font=(CODE_FONT if is_code else H.FONT))
    if attrs.get("underline"):
        r.font.underline = True
    if attrs.get("code") and not code_block:
        _shade(r._element.get_or_add_rPr(), CODE_FILL)
    return r


def _table_borders(tbl):
    tbl_pr = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "808080")
        borders.append(e)
    tbl_pr.append(borders)


_SENTINEL = object()


def _render_table(cell, cells, size):
    """Рисует таблицу редактора (ячейки Quill) вложенной таблицей DOCX.

    cells — [(row_id, [(text, attrs)])]; подряд идущие ячейки с одинаковым
    row_id образуют строку, столбцы задаются порядком ячеек в строке.
    """
    rows = []
    cur = _SENTINEL
    for rid, content in cells:
        if rid != cur:
            rows.append([])
            cur = rid
        rows[-1].append(content)
    ncol = max((len(r) for r in rows), default=0)
    if not ncol:
        return
    tbl = cell.add_table(rows=len(rows), cols=ncol)
    _table_borders(tbl)
    for ri, r in enumerate(rows):
        for ci in range(ncol):
            par = tbl.cell(ri, ci).paragraphs[0]
            par.paragraph_format.space_before = Pt(1)
            par.paragraph_format.space_after = Pt(1)
            content = r[ci] if ci < len(r) else []
            if not content:
                _run(par, "", {}, size)
            for text, attrs in content:
                _run(par, text, attrs, size)


class FigureCounter:
    """Сквозная нумерация рисунков внутри раздела уязвимости."""
    def __init__(self, prefix):
        self.prefix = prefix
        self.n = 0

    def next(self, caption):
        self.n += 1
        num = f"{self.prefix}.{self.n}" if self.prefix else str(self.n)
        return f"Рисунок {num} – {caption}" if caption else f"Рисунок {num}"


def render(cell, value, figures: FigureCounter, image_path, size=11, img_width_mm=115):
    """Добавляет абзацы форматированного текста в ячейку.

    image_path(id) -> путь к файлу изображения или None.
    """
    # cell может быть ячейкой таблицы или телом документа (для разделов без карточек)
    root = getattr(cell, "_element", None)
    if root is None:
        root = cell.element.body
    ops = parse(value)
    if ops is None:
        if value:
            H.add_runs(cell.add_paragraph(), value, size=size)
        return
    line = []          # [(text, attrs)] текущей строки
    list_no = 0        # номер пункта нумерованного списка
    prev_list = None
    tbl_cells = []     # буфер ячеек таблицы: [(row_id, [(text, attrs)])]

    def flush_table():
        if tbl_cells:
            _render_table(cell, list(tbl_cells), size)
            tbl_cells.clear()

    def flush(line_attrs):
        nonlocal list_no, prev_list
        par = cell.add_paragraph()
        par.paragraph_format.space_before = Pt(1)
        par.paragraph_format.space_after = Pt(1)
        code_block = bool(line_attrs.get("code-block"))
        lst = line_attrs.get("list")
        if lst != prev_list or lst != "ordered":
            list_no = 0
        prev_list = lst
        if code_block:
            par.paragraph_format.space_before = Pt(0)
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.left_indent = Mm(2)
            _shade(par._p.get_or_add_pPr(), CODE_FILL)
        elif lst:
            par.paragraph_format.left_indent = Mm(6)
            par.paragraph_format.first_line_indent = Mm(-4)
            if lst == "ordered":
                list_no += 1
                _run(par, f"{list_no}. ", {}, size)
            else:
                _run(par, "• ", {}, size)
        if not line and not code_block:
            par.add_run("")
        for text, attrs in line:
            _run(par, text, attrs, size, code_block)
        line.clear()

    for op in ops:
        ins = op.get("insert")
        attrs = op.get("attributes") or {}
        if isinstance(ins, dict):
            fig = ins.get("figure")
            if not fig:
                continue
            if line:
                flush({})
            flush_table()
            path = image_path(fig.get("id"))
            if not path:
                continue
            p = cell.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            try:
                p.add_run().add_picture(path, width=Mm(img_width_mm))
            except Exception:
                continue
            cap = cell.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.space_after = Pt(6)
            H._set_run(cap.add_run(figures.next((fig.get("caption") or "").strip())), size=10)
            continue
        if not isinstance(ins, str):
            continue
        parts = ins.split("\n")
        for i, part in enumerate(parts):
            if part:
                line.append((part, attrs))
            if i < len(parts) - 1:
                # перевод строки: атрибуты строки (список, блок кода, ячейка таблицы) — у символа \n
                rid = attrs.get("table")
                if rid:
                    tbl_cells.append((rid, list(line)))
                    line.clear()
                else:
                    flush_table()
                    flush(attrs)
    flush_table()
    if line:
        flush({})
    # Quill всегда завершает документ пустой строкой — убираем её,
    # но только если перед ней не таблица (ячейка не должна заканчиваться таблицей)
    last = cell.paragraphs[-1] if cell.paragraphs else None
    if last is not None and len(cell.paragraphs) > 1 and not last.text and not last._p.xpath(".//w:drawing"):
        prev = last._p.getprevious()
        if prev is None or prev.tag != qn("w:tbl"):
            last._p.getparent().remove(last._p)
    # по стандарту OOXML ячейка не может заканчиваться таблицей — добавляем абзац
    if len(root) and root[-1].tag == qn("w:tbl"):
        cell.add_paragraph()


def legacy_images(cell, images, figures: FigureCounter, image_path, img_width_mm=115):
    """Скриншоты старых карточек (загружены отдельным блоком, не в тексте)."""
    for im in images:
        path = image_path(im["id"])
        if not path:
            continue
        p = cell.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        try:
            p.add_run().add_picture(path, width=Mm(img_width_mm))
        except Exception:
            continue
        # номер в старых подписях вводился вручную — нумеруем заново по разделу
        caption = _LEGACY_NUM.sub("", (im.get("caption") or "").strip())
        cap = cell.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(6)
        H._set_run(cap.add_run(figures.next(caption)), size=10)
