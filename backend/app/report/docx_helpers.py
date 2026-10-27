"""Низкоуровневые помощники для оформления DOCX в стиле отчёта РИТ СЕРВИС."""
from docx.shared import Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Times New Roman"
BASE_SIZE = 12          # pt
HEADER_FILL = "F59369"  # оранжевый шапки таблиц


def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_width(cell, mm):
    cell.width = Mm(mm)
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW"); tcPr.append(tcW)
    tcW.set(qn("w:w"), str(int(mm * 56.7)))
    tcW.set(qn("w:type"), "dxa")


def _set_run(run, bold=False, italic=False, size=BASE_SIZE, color=None, font=FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_runs(paragraph, text, **kw):
    """Добавляет текст в параграф, перевод строки -> мягкий перенос."""
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        if i:
            paragraph.add_run().add_break()
        r = paragraph.add_run(line)
        _set_run(r, **kw)
    return paragraph


def fill_cell(cell, text, bold=False, italic=False, size=BASE_SIZE, color=None,
              align=None, bg=None, bullet_list=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    if align:
        p.alignment = align
    if bullet_list is not None:
        for i, item in enumerate(bullet_list):
            para = p if i == 0 else cell.add_paragraph()
            para.paragraph_format.space_after = Pt(2)
            r = para.add_run(f"{i+1}. {item}")
            _set_run(r, size=size)
    else:
        add_runs(p, text, bold=bold, italic=italic, size=size, color=color)
    if bg:
        set_cell_background(cell, bg)


def style_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "808080")
        borders.append(el)
    tblPr.append(borders)


def new_paragraph(doc, text="", style=None, bold=False, italic=False, size=BASE_SIZE,
                  align=None, color=None, space_before=None, space_after=6, first_line=True):
    p = doc.add_paragraph()
    if style:
        try:
            p.style = style
        except KeyError:
            pass
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    if space_before is not None:
        pf.space_before = Pt(space_before)
    if align:
        p.alignment = align
    if first_line and align is None:
        pf.first_line_indent = Mm(10)
    if text:
        add_runs(p, text, bold=bold, italic=italic, size=size, color=color)
    return p


# ---------------------------------------------------------------------------
# Заголовки: встроенные стили Word «Heading 1–3» (нужны для оглавления и
# панели навигации) + закладка _TocN для ссылок и номеров страниц.
# ---------------------------------------------------------------------------
TOC_ENTRIES = []   # [(уровень, текст, имя закладки)] — заполняется при сборке
_bm_id = [0]


def reset_toc():
    TOC_ENTRIES.clear()
    _bm_id[0] = 0


def _bookmark(paragraph, name):
    _bm_id[0] += 1
    bid = str(1000 + _bm_id[0])
    start = OxmlElement("w:bookmarkStart"); start.set(qn("w:id"), bid); start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd"); end.set(qn("w:id"), bid)
    pPr = paragraph._p.pPr
    if pPr is not None:
        pPr.addnext(start)
    else:
        paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _heading(doc, level, full, size, italic=False, before=12, after=6):
    p = doc.add_paragraph(style=f"Heading {level}")
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    pf.keep_with_next = True; pf.first_line_indent = Mm(0)
    r = p.add_run(full)
    _set_run(r, bold=True, italic=italic, size=size, color="000000")
    name = f"_Toc{len(TOC_ENTRIES) + 1:05d}"
    _bookmark(p, name)
    TOC_ENTRIES.append((level, full, name))
    return p


def heading1(doc, text, number=None):
    full = f"{number} {text.upper()}" if number else text.upper()
    return _heading(doc, 1, full, 14, before=18)


def heading2(doc, text, number=None):
    full = f"{number} {text}" if number else text
    return _heading(doc, 2, full, 13)


def heading3(doc, text, number=None):
    full = f"{number} {text}" if number else text
    return _heading(doc, 3, full, 12, italic=True, before=8, after=4)


def setup_heading_styles(doc):
    """Убирает у встроенных заголовков тематические шрифты и цвет."""
    for lvl, size in ((1, 14), (2, 13), (3, 12)):
        st = doc.styles[f"Heading {lvl}"]
        st.font.name = FONT; st.font.size = Pt(size); st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts"); rpr.append(rf)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rf.attrib.pop(qn(attr), None)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(attr), FONT)


# ---------------------------------------------------------------------------
# Оглавление: поле TOC. Результат поля заранее заполнен строками с полями
# PAGEREF, поэтому оглавление видно сразу; Word при открытии обновляет поля
# (updateFields) и пересчитывает номера страниц.
# ---------------------------------------------------------------------------
def _fld(run_parent, kind, instr=None):
    r = OxmlElement("w:r")
    if instr is not None:
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
        r.append(it)
    else:
        fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind)
        r.append(fc)
    run_parent.append(r)


def _text_run(p, text, size=12, bold=False):
    r = p.add_run(text)
    _set_run(r, size=size, bold=bold)
    return r


TEXT_WIDTH_MM = 175


def fill_toc(anchor_paragraph):
    """Вставляет поле оглавления после anchor_paragraph."""
    from docx.text.paragraph import Paragraph
    from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
    entries = TOC_ENTRIES or [(1, "Обновите оглавление (F9)", None)]
    prev = anchor_paragraph
    for i, (lvl, text, bm) in enumerate(entries):
        new_p = OxmlElement("w:p")
        prev._p.addnext(new_p)
        p = Paragraph(new_p, prev._parent)
        pf = p.paragraph_format
        pf.left_indent = Mm(6 * (lvl - 1)); pf.first_line_indent = Mm(0)
        pf.space_after = Pt(2); pf.line_spacing = 1.15
        pf.tab_stops.add_tab_stop(Mm(TEXT_WIDTH_MM), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        if i == 0:
            _fld(p._p, "begin"); _fld(p._p, None, ' TOC \\o "1-3" \\h \\z \\u ')
            _fld(p._p, "separate")
        _text_run(p, text, size=11, bold=(lvl == 1))
        if bm:
            _text_run(p, "\t", size=11)
            _fld(p._p, "begin"); _fld(p._p, None, f" PAGEREF {bm} \\h ")
            _fld(p._p, "separate")
            r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = "1"; r.append(t); p._p.append(r)
            _fld(p._p, "end")
        if i == len(entries) - 1:
            _fld(p._p, "end")
        prev = p
    return prev


def enable_update_fields(doc):
    """Word обновит поля (оглавление, номера страниц) при открытии документа."""
    settings = doc.settings.element
    el = settings.find(qn("w:updateFields"))
    if el is None:
        el = OxmlElement("w:updateFields"); settings.append(el)
    el.set(qn("w:val"), "true")


def caption(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    _set_run(r, size=11)
    return p


def bullet(doc, text, size=BASE_SIZE):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Mm(12)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("– " + text)
    _set_run(r, size=size)
    return p
