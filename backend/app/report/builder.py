"""Сборка итогового отчёта DOCX по данным проекта.

Структура повторяет действующий отчёт и финальный шаблон:
  Титул / Лист согласования
  1 Общие положения (1.1–1.8, модель угроз)
  2 Методология (2.1–2.2 инструментарий)
  3 Краткие итоги (авто-таблицы 3.1, 3.2)
  4..N Разделы по видам работ с карточками уязвимостей
  Приложение 1 — границы
  Приложение 2 — программа и методика (ГОСТ Р 58143-2018, CVSS, методики ФСТЭК)
"""
import io
import os
import json
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from .. import refdata as R
from . import docx_helpers as H

_APP2 = json.load(open(os.path.join(os.path.dirname(__file__), "appendix2.json"), encoding="utf-8"))
C = WD_ALIGN_PARAGRAPH.CENTER
L = WD_ALIGN_PARAGRAPH.LEFT


def _risk_color(level):
    return R.COLORS.get(level, "FFFFFF")


def _add_image_block(doc, images, upload_dir):
    for im in images:
        path = os.path.join(upload_dir, _filename_for(im))
        if not path or not os.path.exists(path):
            continue
        p = doc.add_paragraph(); p.alignment = C
        p.paragraph_format.space_before = Pt(4)
        try:
            run = p.add_run()
            run.add_picture(path, width=Mm(150))
        except Exception:
            continue
        if im.get("caption"):
            cap = doc.add_paragraph(); cap.alignment = C
            cap.paragraph_format.space_after = Pt(6)
            r = cap.add_run(im["caption"]); H._set_run(r, size=11)


# builder работает со словарём image из serializers (url=/api/images/ID),
# но нам нужен файл; передаём отдельную карту id->filename
_IMG_MAP = {}
_UPLOAD_DIR = ""


def _filename_for(im):
    return _IMG_MAP.get(im["id"], "")


# --------------------------------------------------------------------------
def _setup_page(doc, confidential="Конфиденциально"):
    sec = doc.sections[0]
    sec.page_width = Mm(210); sec.page_height = Mm(297)
    sec.top_margin = Mm(20); sec.bottom_margin = Mm(20)
    sec.left_margin = Mm(20); sec.right_margin = Mm(15)
    # нижний колонтитул: "Конфиденциально ... стр. N"
    footer = sec.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.text = ""
    tab_stops = p.paragraph_format.tab_stops
    from docx.enum.text import WD_TAB_ALIGNMENT
    tab_stops.add_tab_stop(Mm(175 / 2), WD_TAB_ALIGNMENT.CENTER)
    tab_stops.add_tab_stop(Mm(175), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run(confidential); H._set_run(r, size=10)
    p.add_run("\t")
    r = p.add_run("стр. "); H._set_run(r, size=10)
    _add_field(p, "PAGE")
    # базовый стиль Normal
    st = doc.styles["Normal"]
    st.font.name = H.FONT; st.font.size = Pt(H.BASE_SIZE)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), H.FONT)
    pf = st.paragraph_format
    pf.line_spacing = 1.3; pf.space_after = Pt(6)


def _add_field(paragraph, field):
    run = paragraph.add_run()
    fldChar1 = OxmlElement("w:fldChar"); fldChar1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = f" {field} "
    fldChar2 = OxmlElement("w:fldChar"); fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1); run._r.append(instr); run._r.append(fldChar2)
    H._set_run(run, size=10)


def _make_table(doc, cols):
    t = doc.add_table(rows=0, cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    H.style_table_borders(t)
    return t


def _header_row(table, titles, widths=None):
    row = table.add_row()
    for i, title in enumerate(titles):
        H.fill_cell(row.cells[i], title, bold=True, align=C, bg=H.HEADER_FILL, size=11)
        if widths:
            H.set_cell_width(row.cells[i], widths[i])
    return row


# ==========================================================================
# ТИТУЛ + ЛИСТ СОГЛАСОВАНИЯ
# ==========================================================================
def _title_page(doc, p, settings):
    for _ in range(2):
        doc.add_paragraph()
    t = doc.add_paragraph(); t.alignment = C
    r = t.add_run("ОТЧЕТ"); H._set_run(r, bold=True, size=16)
    sub = doc.add_paragraph(); sub.alignment = C
    r = sub.add_run("о результатах проведения тестирования объектов информационной инфраструктуры "
                    "на предмет проникновений и анализ уязвимостей информационной безопасности "
                    "объектов информационной инфраструктуры")
    H._set_run(r, bold=True, size=13)
    if p.get("client_name"):
        org = doc.add_paragraph(); org.alignment = C
        r = org.add_run(p["client_name"]); H._set_run(r, bold=True, size=14)
    doc.add_paragraph()
    # лист согласования
    lbl = doc.add_paragraph(); lbl.alignment = C
    r = lbl.add_run("ЛИСТ СОГЛАСОВАНИЯ"); H._set_run(r, bold=True, size=12)
    t = _make_table(doc, 5)
    _header_row(t, ["Наименование организации", "Должность", "Фамилия, имя, отчество", "Подпись", "Дата"],
                widths=[45, 35, 45, 25, 25])
    # подписанты вводятся в проекте; организация по умолчанию — Заказчик / Исполнитель из настроек
    approvals = p.get("approvals") or []
    default_org = {"client": p.get("client_name") or "",
                   "executor": (settings.get("executor") or {}).get("name", "")}
    for side, label in (("client", "Со стороны Заказчика"), ("executor", "Со стороны Исполнителя")):
        row = t.add_row()
        H.fill_cell(row.cells[0].merge(row.cells[4]), label, bold=True, italic=True)
        people = [a for a in approvals if a.get("side") == side] or [{}]
        for a in people:
            row = t.add_row()
            H.fill_cell(row.cells[0], a.get("org") or default_org[side])
            H.fill_cell(row.cells[1], a.get("position", ""))
            H.fill_cell(row.cells[2], a.get("name", ""))
            H.fill_cell(row.cells[3], "")
            H.fill_cell(row.cells[4], a.get("date", ""))
    doc.add_page_break()


# ==========================================================================
# Раздел 1
# ==========================================================================
def _section_general(doc, p, settings, summary):
    H.heading1(doc, "Общие положения", "1")

    H.heading2(doc, "Описание объекта оценки", "1.1")
    methods = p.get("test_methods") or []
    mtext = " и ".join(R.TEST_METHODS[m] for m in methods if m in R.TEST_METHODS)
    desc = p.get("object_description") or ""
    if not desc:
        client = p.get("client_name") or "Заказчика"
        desc = (f"Настоящий Отчет содержит экспертное заключение о результатах услуги по тестированию "
                f"объектов информационной инфраструктуры на предмет проникновений и анализ уязвимостей "
                f"информационной безопасности объектов информационной инфраструктуры {client}"
                + (f" методами {mtext}" if mtext else "") + ".")
    H.new_paragraph(doc, desc, align=L)

    H.heading2(doc, "Общее описание проводимых работ", "1.2")
    H.new_paragraph(doc, "Задачами, решаемыми в рамках оказания услуги, являются:", align=L, first_line=False)
    tasks = [R.WORK_TYPES[w]["task"] for w in p["work_types"] if w in R.WORK_TYPES] + R.DEFAULT_TASKS_TAIL
    for i, task in enumerate(tasks):
        H.bullet(doc, task + (";" if i < len(tasks) - 1 else "."))

    H.heading2(doc, "Технологические участки", "1.3")
    H.new_paragraph(doc, p.get("tech_sites") or R.DEFAULT_TECH_SITES, align=L)

    H.heading2(doc, "Краткие итоги тестирования на проникновение", "1.4")
    H.new_paragraph(doc, "Оценки общего уровня защищенности инфраструктуры Заказчика приведены в таблицах ниже.",
                    align=L)
    n = 1
    for wt in p["work_types"]:
        sec = summary["security"].get(wt)
        if not sec:
            continue
        H.caption(doc, f"Таблица 1.4.{n} – {sec['title']}")
        t = _make_table(doc, 3)
        _header_row(t, ["Преимущества", "Недостатки", "Общий уровень защищенности"], widths=[50, 95, 35])
        row = t.add_row()
        H.fill_cell(row.cells[0], "", bullet_list=sec.get("advantages") or ["—"])
        dis = sec["disadvantages"] or ["Недостатков не выявлено"]
        H.fill_cell(row.cells[1], "", bullet_list=dis)
        H.fill_cell(row.cells[2], sec["level"], bold=True, align=C,
                    bg=R.SECURITY_COLORS.get(sec["level"]))
        n += 1

    H.heading2(doc, "Период проведения тестирования на проникновение", "1.5")
    pf, pt = p.get("period_from"), p.get("period_to")
    if pf or pt:
        H.new_paragraph(doc, f"Период проведения тестирования на проникновение объектов информационной "
                             f"инфраструктуры с {pf or '...'} по {pt or '...'}.", align=L)
    else:
        H.new_paragraph(doc, "Период проведения тестирования на проникновение: с ___ по ___.", align=L)

    H.heading2(doc, "Реквизиты сторон", "1.6")
    t = _make_table(doc, 2)
    _header_row(t, ["ЗАКАЗЧИК", "ИСПОЛНИТЕЛЬ"], widths=[90, 90])
    row = t.add_row()
    H.fill_cell(row.cells[0], (p.get("client_name") or "") + "\n" + (p.get("client_requisites") or ""))
    ex = settings.get("executor", {})
    H.fill_cell(row.cells[1], (ex.get("name", "") + "\n" + ex.get("requisites", "")))

    H.heading2(doc, "Исполнители тестирования на проникновение", "1.7")
    H.caption(doc, "Таблица 1.7.1 – Специалисты со стороны Исполнителя")
    t = _make_table(doc, 3)
    _header_row(t, ["Наименование организации", "Должность", "Фамилия, имя, отчество"], widths=[55, 70, 55])
    execs = p.get("executors") or [{}]
    for e in execs:
        row = t.add_row()
        H.fill_cell(row.cells[0], e.get("org", settings.get("executor", {}).get("name", "")))
        H.fill_cell(row.cells[1], e.get("position", "Специалист по информационной безопасности"))
        H.fill_cell(row.cells[2], e.get("name", ""))

    _threat_model(doc, p)


def _threat_model(doc, p):
    H.heading2(doc, "Модель угроз и нарушителей", "1.8")
    H.new_paragraph(doc,
        "При определении основных типов угроз безопасности информации использованы подходы, установленные "
        "методическим документом «Методика оценки угроз безопасности информации» ФСТЭК России, с учетом "
        "анализа информации об информационных системах и сведений, содержащихся в банке данных угроз "
        "безопасности информации ФСТЭК России.", align=L)
    # в отчёт выводятся только выбранные нарушители (include != False)
    all_intruders = p.get("intruders") or R.DEFAULT_INTRUDERS
    intruders = [it for it in all_intruders if it.get("include", True)] or all_intruders
    H.new_paragraph(doc, "Для исследуемой системы выделены актуальные типы нарушителей:", align=L, first_line=False)
    for i, it in enumerate(intruders):
        H.new_paragraph(doc, f"тип {it['code']} – {it['name']}" + (";" if i < len(intruders)-1 else "."),
                        align=L)

    H.caption(doc, "Таблица 1.8.1 – Информация о выделенных нарушителях")
    t = _make_table(doc, 2)
    _header_row(t, ["Критерий", "Значение"], widths=[50, 130])
    fields = [("Цель (мотив)", "goals"), ("Интерфейсы доступа", "interfaces"),
              ("Ограничения доступа", "restrictions"), ("Технические средства", "tools"),
              ("Навыки", "skills"), ("Вероятность наличия", "probability")]
    for it in intruders:
        row = t.add_row()
        cell = row.cells[0].merge(row.cells[1])
        H.fill_cell(cell, f"Нарушитель {it['code']} – {it['name']}", bold=True, italic=True, bg="FDE3D3")
        for label, key in fields:
            row = t.add_row()
            H.fill_cell(row.cells[0], label)
            H.fill_cell(row.cells[1], it.get(key, ""))

    # только интерфейсы, относящиеся к видам работ проекта
    interfaces = p.get("active_interfaces") or []
    if interfaces:
        H.new_paragraph(doc, "В рамках моделирования угроз определены интерфейсы объекта воздействия – элементы "
                             "исследуемой системы, через которые возможна реализация угроз безопасности информации "
                             "(Таблица 1.8.2).", align=L)
        H.caption(doc, "Таблица 1.8.2 – Интерфейсы воздействия")
        t = _make_table(doc, 5)
        _header_row(t, ["№", "Интерфейс / канал воздействия", "Доступен нарушителям", "Описание", "Примечание"],
                    widths=[15, 45, 25, 55, 40])
        for intf in interfaces:
            row = t.add_row()
            for c, key in zip(row.cells, ["code", "name", "intruders", "description", "note"]):
                H.fill_cell(c, intf.get(key, ""), size=10)

    threats = p.get("auto_threats") or []
    H.new_paragraph(doc, "При определении актуальных угроз использованы сведения из банка данных угроз "
                         "безопасности информации ФСТЭК России (bdu.fstec.ru). Актуальными признаны угрозы, "
                         "реализация которых подтверждена выявленными уязвимостями; для каждой указаны "
                         "задействованные интерфейсы воздействия и типы нарушителей (Таблица 1.8.3).", align=L)
    H.caption(doc, "Таблица 1.8.3 – Актуальные угрозы информационной безопасности")
    t = _make_table(doc, 6)
    _header_row(t, ["№", "Угроза (наименование)", "Идентификатор БДУ", "Интерфейс",
                    "Нарушитель", "Возможные последствия реализации"],
                widths=[14, 55, 20, 18, 18, 55])
    if threats:
        for th in threats:
            row = t.add_row()
            for c, key in zip(row.cells, ["code", "name", "bdu", "interfaces", "intruders", "consequences"]):
                H.fill_cell(c, th.get(key, ""), size=10)
    else:
        row = t.add_row()
        H.fill_cell(row.cells[0].merge(row.cells[5]),
                    "Актуальные угрозы будут сформированы автоматически после того, как в уязвимостях "
                    "будут указаны последствия эксплуатации (H) или реализуемые угрозы БДУ.", italic=True, align=C)


# ==========================================================================
# Раздел 2
# ==========================================================================
def _section_methodology(doc, p, settings):
    H.heading1(doc, "Методология проведения тестирования на проникновение объектов информационной инфраструктуры", "2")
    H.heading2(doc, "Программа и методика тестирования на проникновение", "2.1")
    H.new_paragraph(doc, "Тестирование на проникновение ресурсов Заказчика выполнено в соответствии с "
                         "ГОСТ Р 58143-2018.", align=L)
    if "web" in p["work_types"]:
        H.new_paragraph(doc, "Анализ защищенности веб-приложений дополнительно выполнен с учётом:", align=L,
                        first_line=False)
        H.bullet(doc, "The Open Web Application Security Project (OWASP) Web Security Testing Guide;")
        H.bullet(doc, "The Open Web Application Security Project (OWASP) Top 10.")
    if "social" in p["work_types"]:
        H.new_paragraph(doc, "Социотехническое тестирование проведено по согласованным с Заказчиком "
                             "сценариям в отношении утверждённой целевой группы сотрудников (раздел "
                             f"{p['summary']['section_no']['social']}).", align=L)
    H.new_paragraph(doc, "Методология проведения тестирования на проникновение объектов информационной "
                         "инфраструктуры приведена в Приложении 2.", align=L)
    H.heading2(doc, "Использованный инструментарий тестирования на проникновение объектов информационной инфраструктуры", "2.2")
    H.new_paragraph(doc, "Список инструментов, используемых при тестировании, представлен в таблице 2.2.1.", align=L)
    H.caption(doc, "Таблица 2.2.1 – Список используемых инструментов")
    t = _make_table(doc, 2)
    _header_row(t, ["Цель применения инструментов", "Инструменты"], widths=[80, 100])
    for tool in (p.get("tools") or settings.get("tools") or R.DEFAULT_TOOLS):
        row = t.add_row()
        H.fill_cell(row.cells[0], tool.get("purpose", ""))
        H.fill_cell(row.cells[1], tool.get("tools", ""))


# ==========================================================================
# Раздел 3 — краткие итоги (авто)
# ==========================================================================
def _section_summary(doc, p, summary, settings):
    H.heading1(doc, "Краткие итоги выполненных работ", "3")
    H.new_paragraph(doc, "В ходе проведения работ по тестированию объектов информационной инфраструктуры на "
                         "предмет проникновений и анализу уязвимостей информационной безопасности выявлены "
                         "уязвимости, которые распределены по уровню риска и представлены в таблице 3.1.", align=L)
    H.caption(doc, "Таблица 3.1 – Обнаруженные уязвимости по уровню риска")
    t = _make_table(doc, 3)
    _header_row(t, ["Уязвимость", "Итоговый риск", "Ссылка в отчете"], widths=[120, 30, 30])
    for r_ in summary["rows"]:
        row = t.add_row()
        H.fill_cell(row.cells[0], r_["title"], size=11)
        H.fill_cell(row.cells[1], r_["final_risk"], align=C, bg=_risk_color(r_["final_risk"]), size=11)
        H.fill_cell(row.cells[2], r_["ref"], align=C, size=11)
    if not summary["rows"]:
        row = t.add_row(); H.fill_cell(row.cells[0].merge(row.cells[2]),
                                       "Уязвимости не добавлены", align=C, italic=True)

    if settings.get("report", {}).get("include_risk_stats", True):
        H.new_paragraph(doc, "Распределение выявленных уязвимостей по уровням итогового риска и видам "
                             "работ представлено в таблице 3.2.", align=L)
        H.caption(doc, "Таблица 3.2 – Количество выявленных уязвимостей по уровням риска")
        wts = [wt for wt in p["work_types"] if R.WORK_TYPES.get(wt, {}).get("kind") == "vulns"]
        t = _make_table(doc, len(wts) + 2)
        first_w, total_w = 50, 25
        col_w = (180 - first_w - total_w) / max(len(wts), 1)
        _header_row(t, ["Уровень итогового риска"] + [R.WORK_TYPES[wt]["title"] for wt in wts] + ["Всего"],
                    widths=[first_w] + [col_w] * len(wts) + [total_w])
        by_wt = summary.get("counts_by_wt", {})
        for lvl in ["Высокий", "Средний", "Низкий", "Информативный"]:
            row = t.add_row()
            H.fill_cell(row.cells[0], lvl, bg=_risk_color(lvl))
            for i, wt in enumerate(wts):
                H.fill_cell(row.cells[i + 1], str(by_wt.get(wt, {}).get(lvl, 0)), align=C)
            H.fill_cell(row.cells[-1], str(summary["counts"].get(lvl, 0)), align=C, bold=True)
        row = t.add_row()
        H.fill_cell(row.cells[0], "Итого", bold=True)
        for i, wt in enumerate(wts):
            H.fill_cell(row.cells[i + 1], str(sum(by_wt.get(wt, {}).values())), bold=True, align=C)
        H.fill_cell(row.cells[-1], str(sum(summary["counts"].values())), bold=True, align=C)
        if summary.get("not_rated"):
            n_nr = summary["not_rated"]
            word = "уязвимости" if n_nr % 10 == 1 and n_nr % 100 != 11 else "уязвимостей"
            H.new_paragraph(doc, f"Для {n_nr} {word} итоговый риск не определён (не заполнены критичность "
                                 "или вероятность эксплуатации), в таблицу 3.2 они не включены.", align=L, italic=True)

    if "social" in p["work_types"]:
        st = p.get("social_stats") or {}
        sec = summary["section_no"]["social"]
        H.heading2(doc, "Результаты социотехнического тестирования", "3.3")
        if not st.get("total"):
            H.new_paragraph(doc, "Данные о целевой группе не внесены.", align=L, italic=True)
        else:
            H.new_paragraph(doc, f"Сводные результаты по сценариям приведены в таблице 3.3.1, подробности — "
                                 f"в разделе {sec}.", align=L)
            H.caption(doc, "Таблица 3.3.1 – Сводные результаты социотехнического тестирования")
            _social_stats_table(doc, st)


# ==========================================================================
# Разделы по видам работ + карточки
# ==========================================================================
def _vuln_card(doc, v, upload_dir, include_fstec):
    t = _make_table(doc, 2)
    def kv(label, value, **kw):
        row = t.add_row()
        H.fill_cell(row.cells[0], label, bold=True, size=11)
        H.set_cell_width(row.cells[0], 55)
        H.fill_cell(row.cells[1], value, size=11, **kw)
        H.set_cell_width(row.cells[1], 125)
        return row

    kv("Описание", v.get("description", ""))
    kv("Границы", v.get("scope", ""))
    kv("Инструменты", v.get("tools", ""))
    kv("Время обнаружения", v.get("detected_at", ""))
    score = v.get("cvss_score")
    full = v.get("cvss_vector_full") or v.get("cvss_vector", "")
    cvss_txt = f"{score} ({full})" if score is not None else (full or "—")
    kv("CVSSv3.1", cvss_txt)
    kv("Критичность", v.get("severity", ""), bold=True, bg=_risk_color(v.get("severity")))
    kv("Вероятность эксплуатации", v.get("probability", ""))
    kv("Итоговый риск", v.get("final_risk", ""), bold=True, bg=_risk_color(v.get("final_risk")))
    kv("Модели нарушителя", v.get("intruder_models", ""))
    if v.get("work_type") == "web" and v.get("owasp"):
        owasp = {o["key"]: o["name"] for o in R.OWASP_TOP10}
        kv("Категория OWASP Top 10", owasp.get(v["owasp"], v["owasp"]))

    fstec = v.get("fstec")
    if include_fstec and fstec:
        comp = fstec["components"]
        kv("Тип компонента (K)", f"{comp['K']['name']} (k × K = {comp['K']['kv']})")
        kv("Количество уязвимых компонентов (L)", f"{comp['L']['name']} (l × L = {comp['L']['lv']})")
        kv("Доступность из сети «Интернет» (P)", f"{comp['P']['name']} (p × P = {comp['P']['pv']})")
        kv("Эксплуатация уязвимости (E)", f"{comp['E']['name']} (e × E = {comp['E']['ev']})")
        kv("Последствия эксплуатации (H)", f"{comp['H']['name']} (h × H = {comp['H']['hv']})")
        kv("Уровень критичности по методике ФСТЭК",
           f"V = {fstec['value']}  ({fstec['level']})", bold=True)

    # PoC
    poc_imgs = [im for im in v.get("images", []) if im.get("block") == "poc"]
    row = t.add_row()
    cell = row.cells[0].merge(row.cells[1])
    H.fill_cell(cell, "Proof of Concept", bold=True)
    if v.get("poc"):
        p = cell.add_paragraph(); p.paragraph_format.space_before = Pt(2)
        H.add_runs(p, v["poc"], size=11)
    _cell_images(cell, poc_imgs, upload_dir)

    row = t.add_row()
    cell = row.cells[0].merge(row.cells[1])
    H.fill_cell(cell, "Результаты эксплуатации уязвимости", bold=True)
    if v.get("exploitation_result"):
        p = cell.add_paragraph(); H.add_runs(p, v["exploitation_result"], size=11)
    _cell_images(cell, [im for im in v.get("images", []) if im.get("block") == "result"], upload_dir)

    row = t.add_row()
    cell = row.cells[0].merge(row.cells[1])
    H.fill_cell(cell, "Рекомендации", bold=True)
    if v.get("recommendations"):
        p = cell.add_paragraph(); H.add_runs(p, v["recommendations"], size=11)


def _cell_images(cell, images, upload_dir):
    for im in images:
        fname = _IMG_MAP.get(im["id"], "")
        path = os.path.join(upload_dir, fname) if fname else ""
        if not path or not os.path.exists(path):
            continue
        p = cell.add_paragraph(); p.alignment = C
        try:
            p.add_run().add_picture(path, width=Mm(115))
        except Exception:
            continue
        if im.get("caption"):
            cap = cell.add_paragraph(); cap.alignment = C
            r = cap.add_run(im["caption"]); H._set_run(r, size=10)


def _work_sections(doc, p, summary, settings):
    include_fstec = settings.get("report", {}).get("include_fstec_v", True)
    vulns_by_wt = {}
    for v in p.get("vulns", []):
        vulns_by_wt.setdefault(v["work_type"], []).append(v)
    # сохранить порядок как в таблице 3.1
    order = {}
    for r_ in summary["rows"]:
        order.setdefault(r_["work_type"], []).append(r_["id"])

    for wt in p["work_types"]:
        info = R.WORK_TYPES.get(wt)
        if not info:
            continue
        sec_no = summary["section_no"][wt]
        H.heading1(doc, info["section_title"], str(sec_no))
        if info.get("kind") == "social":
            _social_section(doc, p, sec_no)
            continue
        ids = order.get(wt, [])
        vmap = {v["id"]: v for v in vulns_by_wt.get(wt, [])}
        if not ids:
            H.new_paragraph(doc, "В рамках данного вида работ уязвимости не добавлены.", align=L, italic=True)
            continue
        for i, vid in enumerate(ids, 1):
            v = vmap.get(vid)
            if not v:
                continue
            H.heading2(doc, v.get("title", "Без названия"), f"{sec_no}.{i}")
            _vuln_card(doc, v, settings["_upload_dir"], include_fstec)


# ==========================================================================
# Социотехническое тестирование
# ==========================================================================
def _plural(n, one, few, many):
    if n % 10 == 1 and n % 100 != 11:
        return one
    if 2 <= n % 10 <= 4 and not 12 <= n % 100 <= 14:
        return few
    return many


def _social_stats_table(doc, st):
    t = _make_table(doc, 6)
    _header_row(t, ["Сценарий", "Целевая группа", "Открыли сообщение", "Перешли по ссылке",
                    "Ввели учётные данные", "Сообщили в службу ИБ"], widths=[50, 26, 26, 26, 26, 26])

    def cells(row, d, bold=False):
        total = d.get("total", 0)
        H.fill_cell(row.cells[1], str(total), align=C, size=10, bold=bold)
        for i, key in enumerate(("opened", "clicked", "submitted", "reported"), 2):
            H.fill_cell(row.cells[i], f"{d.get(key, 0)} ({d.get(key + '_pct', 0)}%)", align=C, size=10, bold=bold)

    for d in st.get("by_scenario") or []:
        if not d.get("total"):
            continue
        row = t.add_row()
        H.fill_cell(row.cells[0], d["name"], size=10)
        cells(row, d)
    row = t.add_row()
    H.fill_cell(row.cells[0], "Всего", bold=True, size=10)
    cells(row, st, bold=True)


def _project_images(doc, p, section):
    imgs = sorted([im for im in p.get("images") or [] if im.get("section") == section],
                  key=lambda x: x.get("order_index", 0))
    for im in imgs:
        fname = _IMG_MAP.get(f"p{im['id']}", "")
        path = os.path.join(_UPLOAD_DIR, fname) if fname else ""
        if not path or not os.path.exists(path):
            continue
        par = doc.add_paragraph(); par.alignment = C
        try:
            par.add_run().add_picture(path, width=Mm(150))
        except Exception:
            continue
        if im.get("caption"):
            cap = doc.add_paragraph(); cap.alignment = C
            cap.paragraph_format.space_after = Pt(6)
            H._set_run(cap.add_run(im["caption"]), size=11)


def _social_section(doc, p, sec):
    social = p.get("social") or {}
    st = p.get("social_stats") or {}
    scenarios = social.get("scenarios") or []

    H.heading2(doc, "Общие положения", f"{sec}.1")
    H.new_paragraph(doc, "Целью проведения социотехнического тестирования является оценка осведомленности "
                         "сотрудников Заказчика в области информационной безопасности и их готовности "
                         "противостоять атакам методами социальной инженерии.", align=L)
    H.new_paragraph(doc, "Сценарии и целевая группа согласованы с Заказчиком до начала работ. "
                         + (f"В целевую группу вошли {st['total']} {_plural(st['total'], 'сотрудник', 'сотрудника', 'сотрудников')}; "
                            "перечень приведён в Приложении 1."
                            if st.get("total") else "Перечень целевой группы приведён в Приложении 1."), align=L)

    H.heading2(doc, "Сценарии социальной инженерии", f"{sec}.2")
    if not scenarios:
        H.new_paragraph(doc, "Сценарии не добавлены.", align=L, italic=True)
    else:
        H.caption(doc, f"Таблица {sec}.2.1 – Сценарии социотехнического тестирования")
        t = _make_table(doc, 4)
        _header_row(t, ["№", "Сценарий", "Целевая группа", "Повод / триггер"], widths=[12, 58, 55, 55])
        for i, sc in enumerate(scenarios, 1):
            row = t.add_row()
            H.fill_cell(row.cells[0], str(i), align=C, size=10)
            H.fill_cell(row.cells[1], sc.get("name", ""), size=10)
            H.fill_cell(row.cells[2], sc.get("target_group", ""), size=10)
            H.fill_cell(row.cells[3], sc.get("trigger", ""), size=10)
        for i, sc in enumerate(scenarios, 1):
            H.heading3(doc, sc.get("name") or f"Сценарий {i}", f"{sec}.2.{i}")
            if sc.get("description"):
                H.new_paragraph(doc, sc["description"], align=L)
            _project_images(doc, p, f"social:{sc.get('id')}")

    n = 3
    if (social.get("preparation") or "").strip():
        H.heading2(doc, "Подготовка к проведению социотехнического тестирования", f"{sec}.{n}")
        H.new_paragraph(doc, social["preparation"], align=L)
        n += 1

    H.heading2(doc, "Результаты социотехнического тестирования", f"{sec}.{n}")
    H.new_paragraph(doc, "В ходе тестирования учитывались следующие характеристики:", align=L, first_line=False)
    for i, txt in enumerate(["открывалось ли сообщение сотрудником Заказчика",
                             "переходил ли сотрудник по ссылке на подготовленный ресурс",
                             "вводил ли сотрудник аутентификационные данные",
                             "сообщал ли сотрудник о подозрительном сообщении в службу ИБ"]):
        H.bullet(doc, txt + (";" if i < 3 else "."))
    if st.get("total"):
        H.caption(doc, f"Таблица {sec}.{n}.1 – Сводная информация по результатам социотехнического тестирования")
        _social_stats_table(doc, st)
        if st.get("affected"):
            H.new_paragraph(doc, "Сотрудники, выполнившие целевые действия, приведены в таблице "
                                 f"{sec}.{n}.2.", align=L)
            H.caption(doc, f"Таблица {sec}.{n}.2 – Сотрудники, выполнившие целевые действия")
            t = _make_table(doc, 4)
            _header_row(t, ["№", "Сотрудник / контакт", "Сценарий", "Выполненные действия"], widths=[12, 60, 50, 58])
            for i, a in enumerate(st["affected"], 1):
                row = t.add_row()
                H.fill_cell(row.cells[0], str(i), align=C, size=10)
                H.fill_cell(row.cells[1], a["contact"], size=10)
                H.fill_cell(row.cells[2], a["scenario"], size=10)
                H.fill_cell(row.cells[3], a["actions"], size=10)
        else:
            H.new_paragraph(doc, "Сотрудники целевой группы не переходили по ссылкам и не вводили учётные данные.",
                            align=L)
    else:
        H.new_paragraph(doc, "Данные о целевой группе не внесены.", align=L, italic=True)
    n += 1

    H.heading2(doc, "Рекомендации", f"{sec}.{n}")
    H.new_paragraph(doc, "Для устранения выявленных недостатков рекомендуется:", align=L, first_line=False)
    recs = []
    if st.get("submitted"):
        recs.append("незамедлительно сменить пароли сотрудникам, которые ввели учётные данные (критический приоритет)")
    if st.get("clicked"):
        recs.append("провести внеплановое обучение по противодействию атакам методами социальной инженерии "
                    "для сотрудников, выполнивших целевые действия (высокий приоритет)")
    recs += [x.strip() for x in (social.get("recommendations") or "").split("\n") if x.strip()]
    recs += ["внедрить многофакторную аутентификацию для критичных учётных записей",
             "проводить регулярное обучение и тестирование сотрудников (не реже одного раза в 6 месяцев)"]
    for i, r_ in enumerate(recs):
        H.bullet(doc, r_ + (";" if i < len(recs) - 1 else "."))


# ==========================================================================
# Приложения
# ==========================================================================
def _appendix1(doc, p):
    doc.add_page_break()
    H.heading1(doc, "Приложение 1 – Границы исследования", None)
    H.new_paragraph(doc, "Исходная информация, предоставленная Заказчиком для проведения тестирования на "
                         "проникновение, приведена в таблицах ниже.", align=L)
    scope = p.get("scope") or {}
    tables = [
        ("ips", "Таблица П1.{n} – IP-адреса", ["№", "IP-адрес / диапазон", "Описание"], ["address", "description"]),
        ("fqdns", "Таблица П1.{n} – Доменные имена (FQDN)", ["№", "FQDN", "Описание"], ["name", "description"]),
        ("segments", "Таблица П1.{n} – Сетевые сегменты", ["№", "Сетевой сегмент", "Назначение"], ["segment", "purpose"]),
        ("web_apps", "Таблица П1.{n} – Веб-приложения", ["№", "URL", "Описание"], ["url", "description"]),
    ]
    n = 0
    for key, cap, headers, fields in tables:
        rows = [r for r in (scope.get(key) or []) if any(str(r.get(f, "")).strip() for f in fields)]
        if not rows:
            continue
        n += 1
        H.caption(doc, cap.format(n=n))
        t = _make_table(doc, 3)
        _header_row(t, headers, widths=[12, 78, 90])
        for i, r in enumerate(rows, 1):
            row = t.add_row()
            H.fill_cell(row.cells[0], str(i), align=C, size=10)
            H.fill_cell(row.cells[1], r.get(fields[0], ""), size=10)
            H.fill_cell(row.cells[2], r.get(fields[1], ""), size=10)
    if "social" in p["work_types"]:
        social = p.get("social") or {}
        names = {sc.get("id"): sc.get("name") or f"Сценарий {i}"
                 for i, sc in enumerate(social.get("scenarios") or [], 1)}
        targets = [t for t in (social.get("targets") or []) if str(t.get("contact", "")).strip()]
        if targets:
            n += 1
            H.caption(doc, f"Таблица П1.{n} – Целевая группа социотехнического тестирования")
            t = _make_table(doc, 3)
            _header_row(t, ["№", "Сотрудник / контакт", "Сценарий"], widths=[12, 98, 70])
            for i, tg in enumerate(targets, 1):
                row = t.add_row()
                H.fill_cell(row.cells[0], str(i), align=C, size=10)
                H.fill_cell(row.cells[1], tg.get("contact", ""), size=10)
                H.fill_cell(row.cells[2], names.get(tg.get("scenario"), ""), size=10)
    if not n:
        H.new_paragraph(doc, "Границы исследования не заданы.", align=L, italic=True)


def _appendix2(doc, p, settings):
    doc.add_page_break()
    H.heading1(doc, "Приложение 2 – Программа и методика тестирования на проникновение", None)
    _render_app2_static(doc)
    _render_fstec_criticality(doc, settings)
    _render_probability_matrices(doc)
    _render_risk_matrix(doc)


def _render_app2_static(doc):
    """Текстовые разделы Приложения 2 (ГОСТ Р 58143-2018, CVSS).

    Рендерит всё до раздела «Методика оценки вероятности эксплуатации»,
    то есть включает методику оценки критичности по CVSS из действующего отчёта.
    """
    stop_at = "Методика оценки вероятности эксплуатации найденных уязвимостей"
    for block in _APP2:
        t = block["t"]
        if t in ("h2",) and block["text"] == stop_at:
            break
        if t == "h2":
            H.heading2(doc, block["text"], None)
        elif t == "h3":
            H.heading3(doc, block["text"], None)
        elif t == "bullet":
            H.bullet(doc, block["text"])
        elif t == "caption":
            H.caption(doc, block["text"])
        elif t == "p":
            H.new_paragraph(doc, block["text"], align=L)
        elif t == "table":
            rows = block["rows"]
            tbl = _make_table(doc, len(rows[0]))
            _header_row(tbl, rows[0])
            for r in rows[1:]:
                row = tbl.add_row()
                for i, val in enumerate(r):
                    if i < len(row.cells):
                        H.fill_cell(row.cells[i], val, size=10)


def _render_fstec_criticality(doc, settings):
    H.heading2(doc, "Методика оценки критичности найденных уязвимостей (ФСТЭК России)", None)
    H.new_paragraph(doc,
        "Оценка критичности найденных уязвимостей выполнена с использованием общей системы оценки "
        "уязвимостей CVSS v3.1 (базовые метрики) и методического документа «Методика оценки уровня "
        "критичности уязвимостей программных, программно-аппаратных средств» (утв. ФСТЭК России 30.06.2025).",
        align=L)
    H.new_paragraph(doc, "Итоговый показатель критичности уязвимости рассчитывается по формуле:", align=L)
    f = doc.add_paragraph(); f.alignment = C
    r = f.add_run("V = Iᴄᴠꜱꜱ × Iᵢₙբᵣ × (Iₐₜ + Iᵢₘₚ),   где  Iᵢₙբᵣ = k·K + l·L + p·P,  Iₐₜ = e·E,  Iᵢₘₚ = h·H")
    H._set_run(r, italic=True, size=11)
    weights = (settings.get("fstec") or R.DEFAULT_SETTINGS["fstec"])["weights"]
    H.new_paragraph(doc, f"Весовые коэффициенты: k = {weights['k']}; l = {weights['l']}; p = {weights['p']}; "
                         f"e = {weights['e']}; h = {weights['h']}.", align=L)
    fstec = settings.get("fstec") or R.DEFAULT_SETTINGS["fstec"]
    titles = {"K": "Показатель K – тип компонента", "L": "Показатель L – количество уязвимых компонентов",
              "P": "Показатель P – влияние на защиту периметра", "E": "Показатель E – эксплуатация уязвимости",
              "H": "Показатель H – возможные последствия эксплуатации"}
    w_for = {"K": "k", "L": "l", "P": "p", "E": "e", "H": "h"}
    for key in ("K", "L", "P", "E", "H"):
        H.caption(doc, titles[key])
        tbl = _make_table(doc, 3)
        _header_row(tbl, ["Наименование", "Значение", f"{w_for[key]} × {key}"], widths=[120, 30, 30])
        for rowd in fstec[key]:
            row = tbl.add_row()
            H.fill_cell(row.cells[0], rowd["name"], size=10)
            H.fill_cell(row.cells[1], str(rowd["value"]), align=C, size=10)
            H.fill_cell(row.cells[2], str(round(weights[w_for[key]] * rowd["value"], 3)), align=C, size=10)
    H.caption(doc, "Уровни критичности уязвимости")
    tbl = _make_table(doc, 2)
    _header_row(tbl, ["Уровень критичности", "Значение V"], widths=[90, 90])
    for cond, lvl in [("V > 8,0", "Критический"), ("5,0 ≤ V ≤ 8,0", "Высокий"),
                      ("2,0 ≤ V < 5,0", "Средний"), ("V < 2,0", "Низкий")]:
        row = tbl.add_row()
        H.fill_cell(row.cells[0], lvl, bold=True)
        H.fill_cell(row.cells[1], cond, align=C)


def _render_probability_matrices(doc):
    for block in _APP2:
        if block["t"] == "h2" and "вероятности эксплуатации" in block["text"]:
            idx = _APP2.index(block)
            break
    else:
        return
    # отрисовать от этого h2 до раздела риска
    i = idx
    while i < len(_APP2):
        block = _APP2[i]
        t = block["t"]
        if t == "h2" and "итогового уровня риска" in block["text"]:
            break
        if t == "h2":
            H.heading2(doc, block["text"], None)
        elif t == "caption":
            H.caption(doc, block["text"])
        elif t == "p":
            H.new_paragraph(doc, block["text"], align=L)
        elif t == "matrix" and block.get("kind") == "probability":
            _draw_matrix(doc, "Доступность", "Простота эксплуатации", R.PROBABILITY_MATRIX)
        elif t == "table":
            rows = block["rows"]
            tbl = _make_table(doc, len(rows[0]))
            _header_row(tbl, rows[0])
            for r in rows[1:]:
                row = tbl.add_row()
                for j, val in enumerate(r):
                    if j < len(row.cells):
                        H.fill_cell(row.cells[j], val, size=10)
        i += 1


def _render_risk_matrix(doc):
    H.heading2(doc, "Методика оценки итогового уровня риска найденных уязвимостей", None)
    H.new_paragraph(doc, "Риск уязвимости рассчитывается на основе критичности уязвимости и вероятности "
                         "эксплуатации уязвимости по таблице ниже.", align=L)
    _draw_matrix(doc, "Критичность уязвимости", "Вероятность эксплуатации", R.RISK_MATRIX, color=True)


def _draw_matrix(doc, row_label, col_label, matrix, color=False):
    cols = R.LEVELS3
    tbl = _make_table(doc, len(cols) + 2)
    hdr = tbl.add_row()
    c0 = hdr.cells[0].merge(hdr.cells[1])
    H.fill_cell(c0, row_label + " \\ " + col_label, bold=True, align=C, bg=H.HEADER_FILL, size=10)
    for i, c in enumerate(cols):
        H.fill_cell(hdr.cells[i + 2], c, bold=True, align=C, bg=H.HEADER_FILL, size=10)
    row_keys = list(matrix.keys())
    for rk in row_keys:
        row = tbl.add_row()
        merged = row.cells[0].merge(row.cells[1])
        H.fill_cell(merged, rk, bold=True, size=10)
        for i, ck in enumerate(cols):
            val = matrix[rk].get(ck, "")
            bg = _risk_color(val) if color else None
            H.fill_cell(row.cells[i + 2], val, align=C, size=10, bg=bg)


# ==========================================================================
# Точка входа
# ==========================================================================
def build_report(project, settings, upload_dir, img_map=None):
    global _IMG_MAP
    _IMG_MAP = img_map or {}
    settings = dict(settings)
    settings["_upload_dir"] = upload_dir
    global _UPLOAD_DIR
    _UPLOAD_DIR = upload_dir

    doc = Document()
    H.reset_toc()
    _setup_page(doc, settings.get("report", {}).get("confidential_label", "Конфиденциально"))
    H.setup_heading_styles(doc)
    summary = project["summary"]

    _title_page(doc, project, settings)
    # оглавление: заголовок страницы, поле TOC вставляется после сборки документа
    toc_title = doc.add_paragraph()
    toc_title.paragraph_format.first_line_indent = Mm(0)
    toc_title.paragraph_format.space_after = Pt(10)
    H._set_run(toc_title.add_run("Оглавление"), bold=True, size=14)
    doc.add_page_break()
    _section_general(doc, project, settings, summary)
    _section_methodology(doc, project, settings)
    _section_summary(doc, project, summary, settings)
    _work_sections(doc, project, summary, settings)
    _appendix1(doc, project)
    _appendix2(doc, project, settings)

    H.fill_toc(toc_title)
    H.enable_update_fields(doc)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
